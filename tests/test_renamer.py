from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from helpers import BASE  # noqa: F401  让 core 可导入

from core.renamer import (
    MAX_NAME,
    RenameItem,
    apply_plan,
    extract_title,
    log_path,
    plan_folder,
    summarize_name,
    undo_last,
)

HEADER_DOC = """广东运峰电力安装有限公司文件
运峰综〔2022］7号
关于成立某市220千伏变电站扩建第三台
主变及其配套线路工程项目部暨任命王敬
同志为项目经理的通知
各部门：
为更好地推进某市220千伏变电站扩建第三台主变及其配套线路工程顺利实施，经公司决议成立项目部。
"""

TABLE_DOC = """表B.0.7
施工单位体系认证报审/报验表
工程名称：某市220千伏变电站扩建第三台主变及其配套线路工程
编号：SD-030600WB23230001-000
"""

WRAPPED_DOC = """广东运峰电力安装有限公司
关于成立某市220千伏变电站扩建第三台主变
及其配套线路工程项目部及启用项目部印章的
图
某省级电网公司某市供电局：
由于工程施工需要，我公司决定成立某市220千伏变电站扩建第三台主变及其配套线路工程项目部。
"""

# 南网《电网建设施工作业指导书》封面：抬头对每一份都一样，题目在“第X部分：”后面
GUIDE_COVER = """中国南方电网
中国南方电网有限责任公司电网建设施工作业指导书
第1部分：输电线路
外拉线悬浮抱杆分解组立铁塔
作业指导书
编码：SDXL-ZW-07
"""

GUIDE_COVER_WITH_WORD = """中国南方电网有限责任公司电网建设施工作业指导书
中国南方电网
第1部分：输电线路
导地线压接作业指导书
编码：SDXL-ZW-10
"""

# 南网质量验收记录表：档号 + 表号 + 题目两行 + 工程编号行
FORM_DOC = """【S005】
Q/CSG表1-B.8灌注桩基础
分项工程质量验收记录(线基8)
工程编号：SD-030600WB23230001
Q/CSG表1-B.8
设计桩号
杆塔型
"""

FORM_DOC_WITH_COMMA = """【S005】
Q/CSG表1-B.9灌注桩基础承台、连梁
分项工程质量验收记录表（线基9)
工程编号：SD-030600WB23230001
Q/CSG表1-B.9
杆塔型
"""

# 图纸：正文开头是图面注记，真图名在右下角图签里
DRAWING = """玉永甲#38
2A1
原线路增挂导线
新建双回线路（常规段）
平面图
平断面图4（2A1~3A23）
广东省建设工程勘察设计出图专用章
单位名称：某市电力设计院有限公司
"""


class TitleExtractionTests(unittest.TestCase):
    def test_red_header_title_merges_wrapped_lines(self):
        self.assertEqual(
            extract_title(HEADER_DOC),
            "关于成立某市220千伏变电站扩建第三台主变及其配套线路工程项目部暨任命王敬同志为项目经理的通知",
        )

    def test_table_number_is_stripped(self):
        title = extract_title("表B.0.19施工质量验收及评定项目划分报审表★\n工程名称：某工程")
        self.assertEqual(title, "施工质量验收及评定项目划分报审表")

    def test_table_number_on_its_own_line(self):
        self.assertEqual(extract_title(TABLE_DOC), "施工单位体系认证报审/报验表")

    def test_field_lines_are_skipped(self):
        title = extract_title("工程名称：某市220千伏变电站扩建第三台主变及其配套线路工程\n安全技术交底")
        self.assertEqual(title, "安全技术交底")


class SummarizeTests(unittest.TestCase):
    def test_project_name_is_removed(self):
        name = summarize_name(HEADER_DOC)
        self.assertEqual(name, "成立项目部暨任命王敬同志为项目经理的通知")

    def test_wrapped_title_and_ocr_dropping(self):
        # 末尾那个“图”是 OCR 认错的“函”，不接进标题
        name = summarize_name(WRAPPED_DOC)
        self.assertEqual(name, "成立项目部及启用项目部印章")

    def test_slash_between_chinese_is_dropped(self):
        name = summarize_name(TABLE_DOC)
        self.assertEqual(name, "施工单位体系认证报审报验表")

    def test_name_is_shortened_to_limit(self):
        long_title = "关于" + "某某某某" * 20 + "工程验收记录"
        name = summarize_name(long_title)
        self.assertLessEqual(len(name), MAX_NAME)
        self.assertTrue(name)

    def test_no_title_returns_empty(self):
        self.assertEqual(summarize_name("12345\n2024年1月1日"), "")

    def test_series_head_is_dropped(self):
        # 抬头是整套书的书名，“作业指导书”是文种，要补在题目后面
        self.assertEqual(
            summarize_name(GUIDE_COVER),
            "输电线路外拉线悬浮抱杆分解组立铁塔作业指导书",
        )

    def test_series_head_doc_word_is_not_doubled(self):
        self.assertEqual(
            summarize_name(GUIDE_COVER_WITH_WORD),
            "输电线路导地线压接作业指导书",
        )

    def test_drawing_uses_name_from_stamp(self):
        self.assertEqual(summarize_name(DRAWING), "平断面图4(2A1~3A23)")

    def test_form_keeps_title_drops_stamp_and_number(self):
        # 表号只留 B.8，档号【S005】和“工程编号：…”都不是文件名
        self.assertEqual(
            summarize_name(FORM_DOC), "B.8灌注桩基础分项工程质量验收记录(线基8)"
        )

    def test_form_title_with_comma_is_not_truncated(self):
        self.assertEqual(
            summarize_name(FORM_DOC_WITH_COMMA),
            "B.9灌注桩基础承台、连梁分项工程质量验收记录表(线基9)",
        )


class ApplyAndUndoTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.src = self.root / "ok.pdf"
        self.src.write_bytes(b"%PDF-1.4 test")

    def tearDown(self):
        self.tmp.cleanup()

    def test_rename_then_undo(self):
        item = RenameItem(path=self.src, old_name=self.src.name, new_name="安全技术交底", ext=".pdf")
        log = log_path(self.root, base=self.root / "logs")
        result = apply_plan([item], log_file=log)
        self.assertEqual(len(result["done"]), 1)
        renamed = self.root / "安全技术交底.pdf"
        self.assertTrue(renamed.exists())
        self.assertFalse(self.src.exists())

        back = undo_last(log)
        self.assertEqual(len(back["done"]), 1)
        self.assertTrue(self.src.exists())
        self.assertFalse(renamed.exists())

    def test_existing_target_is_not_overwritten(self):
        occupied = self.root / "安全技术交底.pdf"
        occupied.write_bytes("已经有一份".encode("utf-8"))
        item = RenameItem(path=self.src, old_name=self.src.name, new_name="安全技术交底", ext=".pdf")
        result = apply_plan([item], log_file=log_path(self.root, base=self.root / "logs"))
        self.assertEqual(len(result["done"]), 0)
        self.assertEqual(len(result["failed"]), 1)
        self.assertEqual(occupied.read_bytes(), "已经有一份".encode("utf-8"))

    def test_plan_folder_names_docx(self):
        from docx import Document

        doc = Document()
        doc.add_paragraph("广东运峰电力安装有限公司")
        doc.add_paragraph("工程质量终身责任承诺书")
        doc.add_paragraph("工程名称：某市220千伏变电站扩建第三台主变及其配套线路工程")
        doc.save(str(self.root / "abc.docx"))
        items = plan_folder(self.root, use_ocr=False)
        by_ext = {item.ext: item for item in items}
        self.assertEqual(by_ext[".docx"].new_name, "工程质量终身责任承诺书")
        self.assertEqual(by_ext[".docx"].source, "Word")


if __name__ == "__main__":
    unittest.main()
