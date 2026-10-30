from __future__ import annotations

import unittest
from io import BytesIO

from helpers import make_checkpoint

from openpyxl import load_workbook

from core.engine import build_stats
from core.exporter import export_report
from core.models import (
    STATE_MATCHED,
    STATE_MISSING,
    STATE_NA,
    STATE_SUSPECT,
    CategoryGap,
    CheckReport,
    Evidence,
    MatchResult,
)


def result(cp, state, score=0.0, reason="", evidence=None):
    return MatchResult(
        checkpoint=cp,
        state=state,
        score=score,
        evidence=evidence or [],
        reason=reason,
        auto_state=state,
        auto_reason=reason,
    )


def make_report(missing_text="水泥出厂合格证"):
    matched_cp = make_checkpoint("混凝土试块报告及报审", seq=1)
    suspect_cp = make_checkpoint("钢筋出厂合格证", seq=2)
    missing_cp = make_checkpoint(missing_text, seq=3)
    na_cp = make_checkpoint("其他文件", optional=True, seq=4)
    results = [
        result(
            matched_cp, STATE_MATCHED, 0.9, "命中",
            [Evidence("8232/混凝土.pdf", 0.9)],
        ),
        result(
            suspect_cp, STATE_SUSPECT, 0.5, "疑似",
            [Evidence("8232/钢筋.pdf", 0.5)],
        ),
        result(missing_cp, STATE_MISSING, 0.0, "未发现匹配文件"),
        result(na_cp, STATE_NA, 0.0, "其他/备查类条目"),
    ]
    stats = build_stats(results)
    return CheckReport(
        meta={
            "project_name": "测试工程",
            "doc_type_name": "基本建设",
            "project_type_name": "82 交流输电线路",
            "unit": "施工",
            "root": r"D:\测试",
            "template_name": "测试模版",
            "template_hash": "abcdef1234567890",
            "generated_at": "2026-10-28T10:00:00",
            "filtered_total": 4,
            "match_threshold": 0.62,
            "suspect_threshold": 0.40,
        },
        results=results,
        category_gaps=[
            CategoryGap(code="820", category="项目准备", path=["基本建设", "项目准备"], count=3)
        ],
        orphan_files=[
            {
                "rel_path": "其他/未归类.pdf",
                "ext": ".pdf",
                "size": 2048,
                "best_score": 0.1,
                "best_cp_id": "CP-8232-0001",
                "best_cp_text": "混凝土试块报告及报审",
            }
        ],
        stats=stats,
    )


class ExporterTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.report = make_report()
        cls.bytes = export_report(
            cls.report, project_name="测试工程", inspector="张三", reviewer="李四", due_days=7
        )
        cls.wb = load_workbook(BytesIO(cls.bytes))

    def test_sheet_names(self):
        self.assertEqual(
            self.wb.sheetnames,
            ["封面", "整改单", "疑似清单", "已匹配", "印章核查", "未归类文件", "统计", "说明"],
        )

    def test_cover_title(self):
        self.assertEqual(self.wb["封面"]["A1"].value, "竣工资料完整性检查整改单")

    def test_rectification_rows(self):
        ws = self.wb["整改单"]
        self.assertEqual(ws["B2"].value, "820")
        self.assertIn("整类缺失", ws["E2"].value)
        self.assertEqual(ws["E3"].value, "水泥出厂合格证")
        self.assertEqual(ws["I3"].value, STATE_MISSING)
        self.assertEqual(ws["A1"].fill.fill_type, "solid")
        self.assertTrue(ws["A1"].font.bold)
        self.assertIsNotNone(ws.freeze_panes)
        self.assertIsNotNone(ws.auto_filter.ref)

    def test_suspect_and_matched_sheets(self):
        self.assertEqual(self.wb["疑似清单"]["E2"].value, "钢筋出厂合格证")
        self.assertEqual(self.wb["已匹配"]["E2"].value, "混凝土试块报告及报审")
        self.assertEqual(self.wb["已匹配"]["F2"].value, STATE_MATCHED)

    def test_orphan_sheet(self):
        self.assertEqual(self.wb["未归类文件"]["B2"].value, "其他/未归类.pdf")

    def test_stats_sheet_contains_totals(self):
        ws = self.wb["统计"]
        values = [str(cell.value) for row in ws.iter_rows() for cell in row if cell.value]
        self.assertIn("资料完整率", values)
        self.assertIn("按三级分类统计", " ".join(values))

    def test_formula_injection_is_escaped(self):
        report = make_report(missing_text="=HYPERLINK(\"http://x\")")
        data = export_report(report, project_name="测试工程")
        ws = load_workbook(BytesIO(data))["整改单"]
        self.assertTrue(str(ws["E3"].value).startswith("'="))

    def test_landscape_page_setup(self):
        self.assertEqual(self.wb["整改单"].page_setup.orientation, "landscape")
        self.assertEqual(self.wb["封面"].page_setup.orientation, "portrait")


if __name__ == "__main__":
    unittest.main()
