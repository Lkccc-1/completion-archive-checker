from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from helpers import BASE

from docx import Document

from core.docx_parser import TemplateParseError, parse_docx
from core.engine import filter_checkpoints

INVENTORY = BASE / "项目A（220kV输电线路工程）PDF文件清单（施工）.docx"


def build_inventory(path: Path, rows: list[tuple[str, str, str]]) -> None:
    doc = Document()
    table = doc.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ("序号", "文件名", "立卷")):
        cell.text = text
    for seq, name, filing in rows:
        cells = table.add_row().cells
        cells[0].text = seq
        cells[1].text = name
        cells[2].text = filing
    doc.save(str(path))


class DocxInventoryTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / "清单.docx"

    def tearDown(self):
        self.tmp.cleanup()

    def test_folder_rows_and_code_chain(self):
        build_inventory(
            self.path,
            [
                ("AX-8230-1 成立文件、开工报告", "AX-8230-1 成立文件、开工报告", "AX-8230-1 成立文件、开工报告"),
                ("1", "关于成立项目部的通知", "施工"),
                ("2", "开工报告及报审表", "施工"),
                ("AX-825-1 试验报告", "AX-825-1 试验报告", "AX-825-1 试验报告"),
                ("1", "线路参数测试试验报告", "施工"),
            ],
        )
        cs = parse_docx(self.path)
        self.assertEqual(cs.meta["mode"], "inventory")
        self.assertEqual(len(cs.checkpoints), 3)
        self.assertEqual([cp.code for cp in cs.checkpoints], ["8230", "8230", "825"])
        self.assertEqual(cs.checkpoints[0].path, ["基本建设", "交流输电线路", "成立文件、开工报告"])
        self.assertEqual(cs.checkpoints[0].seq, 1)
        self.assertEqual(cs.checkpoints[0].filing, "施工")
        self.assertEqual(cs.checkpoints[0].filing_units, ["施工"])
        # 清单里的序号是按卷重排的，检查点 id 仍要唯一
        self.assertEqual(len({cp.id for cp in cs.checkpoints}), 3)
        self.assertEqual(cs.project_types()[0].code, "82")

    def test_filter_by_unit_and_code(self):
        build_inventory(
            self.path,
            [
                ("AX-8230-1 成立文件", "AX-8230-1 成立文件", "AX-8230-1 成立文件"),
                ("1", "开工报告", "施工"),
                ("2", "监理规划", "监理"),
                ("AX-9201-1 设备开箱", "AX-9201-1 设备开箱", "AX-9201-1 设备开箱"),
                ("3", "开箱检查记录", "建设"),
            ],
        )
        cs = parse_docx(self.path)
        picked = filter_checkpoints(cs.checkpoints, "8", "82", "施工")
        self.assertEqual([cp.text for cp in picked], ["开工报告"])

    def test_heading_without_code_warns(self):
        build_inventory(
            self.path,
            [
                ("AX-8230-1 成立文件", "AX-8230-1 成立文件", "AX-8230-1 成立文件"),
                ("1", "开工报告", "施工"),
                ("补充材料", "补充材料", "补充材料"),
                ("2", "补充说明", "施工"),
            ],
        )
        cs = parse_docx(self.path)
        self.assertEqual([cp.code for cp in cs.checkpoints], ["8230", "8230"])
        self.assertTrue(any("没有 3~4 位分类号" in text for text in cs.warnings))

    def test_document_without_known_table_raises(self):
        doc = Document()
        doc.add_table(rows=1, cols=2).rows[0].cells[0].text = "编号"
        doc.save(str(self.path))
        with self.assertRaises(TemplateParseError):
            parse_docx(self.path)

    @unittest.skipUnless(INVENTORY.exists(), "样例文件清单不在仓库里")
    def test_real_inventory_sample(self):
        cs = parse_docx(INVENTORY)
        self.assertEqual(len(cs.checkpoints), 669)
        self.assertEqual(cs.meta["table_rows"], 741)
        self.assertEqual(cs.meta["doc_types"], ["8", "9"])
        self.assertEqual({cp.filing for cp in cs.checkpoints}, {"施工"})


if __name__ == "__main__":
    unittest.main()
