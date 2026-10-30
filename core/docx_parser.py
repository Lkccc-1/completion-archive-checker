"""解析《电网建设项目、设备仪器类分类、归档范围及保管期限表》docx。

模版是一张 12 列、两行表头的大表：
分类号 | 类目名称 | 归档内容 | 主要项目文件 | 执行标准 | 来源 | 立卷 |
建设 | 设计 | 监理 | 施工 | 备注

解析时以“行序 + 编码前缀”共同确定层级：模版里存在少量编码错位
（如 826 下的 8464），单纯按前缀会挂到 846 变电站监理下面，必须按行序纠正。
"""
from __future__ import annotations

import hashlib
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document

from .models import DOC_TYPES, Checkpoint, CheckpointSet, Node, UNITS
from .text import (
    clean_display,
    extract_codes,
    extract_terms,
    is_optional_text,
    join_lines,
    parse_seq,
    split_bullets,
)

COLUMNS = {
    "code": 0,
    "name": 1,
    "archive": 2,
    "files": 3,
    "standard": 4,
    "source": 5,
    "filing": 6,
    "建设": 7,
    "设计": 8,
    "监理": 9,
    "施工": 10,
    "note": 11,
}


class TemplateParseError(ValueError):
    """模版无法解析时抛出，UI 直接展示这条信息。"""


def _find_table(doc: Any) -> tuple[Any, int]:
    for table in doc.tables:
        for idx, row in enumerate(table.rows[:5]):
            text = "".join(cell.text for cell in row.cells)
            if "分类号" in text and "主要项目文件" in text:
                return table, idx
    raise TemplateParseError(
        "未找到包含“分类号”和“主要项目文件”的表，也没有找到“文件名＋立卷”的文件清单表。"
        "请确认上传的是《电网建设项目、设备仪器类分类、归档范围及保管期限表》，"
        "或按案卷列出的《竣工资料 PDF 文件清单》，或改用 YAML 模版。"
    )


def _find_inventory_table(doc: Any) -> tuple[Any, int] | None:
    """找“序号｜文件名｜立卷”式的文件清单表。

    与归档范围表的区别：清单没有“主要项目文件”，只有案卷（合并单元格的标题行）
    加一份份实际文件。用它当检查标准时，核对的是“清单列的每一份在目录里有没有”。
    """
    for table in doc.tables:
        for idx, row in enumerate(table.rows[:8]):
            cells = {cell.text.strip() for cell in row.cells}
            if "文件名" in cells and ("立卷" in cells or "序号" in cells):
                return table, idx
    return None


PROJECT_TYPE_NAMES = {
    "80": "综合",
    "81": "调度自动化、通信",
    "82": "交流输电线路",
    "83": "电力电缆线路",
    "84": "变电站",
    "85": "直流输电线路（含接地极线路）",
    "86": "换流站",
    "89": "其他工程",
    "90": "综合",
    "91": "调度自动化、通信",
    "92": "交流输电线路",
    "93": "电力电缆线路",
    "94": "变电站",
    "95": "直流输电线路（含接地极线路）",
    "96": "换流站",
    "99": "其他工程",
}


def _ensure_chain(
    nodes: dict[str, Node], code: str, leaf_name: str, order: int
) -> str | None:
    """按编码补齐上级节点（8 → 82 → 825/8230），返回叶子节点 key。

    文件清单里只有 3~4 位编码，没有上级类目名。一级（8/9）与二级（82 等）
    用固定名称补齐，三级以下用案卷标题当名称。
    """
    if not code or not code.isdigit():
        return None
    parent_key: str | None = None
    sizes = sorted({size for size in (1, 2, len(code)) if size <= len(code)})
    for size in sizes:
        prefix = code[:size]
        name = DOC_TYPES.get(prefix) or PROJECT_TYPE_NAMES.get(prefix)
        if name is None:
            name = leaf_name or code
        key = f"{parent_key or ''}|{prefix}|{name}"
        if key not in nodes:
            nodes[key] = Node(
                key=key,
                code=prefix,
                name=name,
                level=size,
                parent_key=parent_key,
                order=order,
            )
        parent_key = key
    return parent_key


def _parse_inventory(
    doc: Any, path: Path, table: Any, header_idx: int
) -> "CheckpointSet":
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    header = rows[header_idx]
    idx_name = header.index("文件名")
    idx_seq = header.index("序号") if "序号" in header else -1
    idx_filing = header.index("立卷") if "立卷" in header else -1

    nodes: dict[str, Node] = {}
    checkpoints: list[Checkpoint] = []
    warnings: list[str] = []
    node_order = 0
    folder_raw = ""
    folder_label = ""
    leaf_key: str | None = None
    leaf_code = ""
    skipped = 0

    for row_index in range(header_idx + 1, len(rows)):
        row = rows[row_index]
        if not any(row):
            continue
        cells = row + [""] * (max(idx_name, idx_filing, idx_seq) + 1 - len(row))
        # 合并单元格的案卷标题行：各列文字相同
        if len(set(cell for cell in cells if cell)) == 1:
            folder_raw = cells[0].strip()
            if not folder_raw:
                continue
            codes = extract_codes(folder_raw)
            parts = folder_raw.split(None, 1)
            label = parts[1].strip() if len(parts) > 1 else folder_raw
            if not codes:
                warnings.append(
                    f"第 {row_index + 1} 行案卷“{folder_raw}”没有 3~4 位分类号，"
                    "该卷下的文件按上一卷归档，请核对清单"
                )
                continue
            leaf_code = codes[0]
            folder_label = label
            leaf_key = _ensure_chain(nodes, leaf_code, folder_label, node_order)
            node_order = len(nodes)
            continue

        name = cells[idx_name].strip()
        if not name:
            skipped += 1
            continue
        if leaf_key is None:
            warnings.append(f"第 {row_index + 1} 行文件“{name[:30]}”在案卷标题之前，已跳过")
            skipped += 1
            continue

        raw_seq = cells[idx_seq].strip() if idx_seq >= 0 else ""
        seq = int(raw_seq) if raw_seq.isdigit() else len(checkpoints) + 1
        filing = cells[idx_filing].strip() if idx_filing >= 0 else ""
        checkpoints.append(
            Checkpoint(
                id=f"CP-{leaf_code}-{len(checkpoints) + 1:04d}",
                node_key=leaf_key,
                code=leaf_code,
                category=folder_label,
                path=_build_path(nodes, leaf_key),
                archive_content=folder_label,
                text=name,
                seq=seq,
                source="文件清单",
                filing=filing,
                retention={},
                keywords=extract_terms(name),
                extra={"row": row_index + 1, "folder": folder_raw},
            )
        )

    if not checkpoints:
        raise TemplateParseError(
            "文件清单表识别成功，但没有解析到任何文件行，请检查“文件名”列。"
        )

    title = (doc.core_properties.title or "").strip()
    if not title:
        for paragraph in doc.paragraphs[:3]:
            if paragraph.text.strip():
                title = paragraph.text.strip()
                break
    title = title or path.stem

    warnings.insert(
        0,
        "按“文件清单”模式解析：每一份实际文件是一条检查点，核对的是清单与目录是否一致，"
        "不是归档范围是否齐全。",
    )
    warnings.append(
        f"数据行：{len(rows) - header_idx - 1} 行；解析文件 {len(checkpoints)} 条；"
        f"跳过空行/无归属行 {skipped} 行。"
    )

    return CheckpointSet(
        meta={
            "schema_version": 2,
            "source_file": path.name,
            "source_hash": _sha256(path),
            "template_name": title,
            "parsed_at": datetime.now().isoformat(timespec="seconds"),
            "mode": "inventory",
            "doc_types": sorted({cp.code[0] for cp in checkpoints if cp.code}),
            "project_types": [
                {"code": node.code, "name": node.name}
                for node in sorted(nodes.values(), key=lambda item: item.order)
                if node.level == 2
            ],
            "total_checkpoints": len(checkpoints),
            "total_nodes": len(nodes),
            "table_rows": len(rows),
            "header_row": header_idx + 1,
            "columns": header,
        },
        nodes=sorted(nodes.values(), key=lambda item: item.order),
        checkpoints=checkpoints,
        warnings=warnings,
    )


def _validate_columns(rows: list[list[str]], header_idx: int) -> list[str]:
    if not rows or len(rows[0]) < 12:
        raise TemplateParseError(f"表头列数不足：期望 12 列，实际 {len(rows[0]) if rows else 0} 列")
    header = " | ".join(cell.strip() for cell in rows[header_idx])
    required = ("归档内容", "主要项目文件", "来源", "立卷", "建设", "设计", "监理", "施工")
    missing = [name for name in required if name not in header]
    if missing:
        raise TemplateParseError(
            f"表头缺少列：{'、'.join(missing)}。检测到的表头是：{header[:300]}"
        )
    return [cell.strip() for cell in rows[header_idx]]


def _cell(row: list[str], key: str) -> str:
    index = COLUMNS[key]
    return row[index].strip() if index < len(row) else ""


def _build_path(nodes: dict[str, Node], key: str | None) -> list[str]:
    names: list[str] = []
    seen: set[str] = set()
    while key and key not in seen:
        seen.add(key)
        node = nodes.get(key)
        if node is None:
            break
        names.append(node.name)
        key = node.parent_key
    return list(reversed(names))


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def parse_docx(path: str | Path) -> CheckpointSet:
    path = Path(path)
    try:
        doc = Document(str(path))
    except Exception as exc:  # noqa: BLE001
        raise TemplateParseError(f"无法打开 docx：{exc}") from exc

    try:
        table, header_idx = _find_table(doc)
    except TemplateParseError:
        # 归档范围表找不到时，再试“文件清单”格式：按案卷列实际文件的 docx。
        inventory = _find_inventory_table(doc)
        if inventory is None:
            raise
        return _parse_inventory(doc, path, *inventory)
    rows = [[cell.text for cell in row.cells] for row in table.rows]
    header = _validate_columns(rows, header_idx)

    nodes: dict[str, Node] = {}
    node_order = 0
    current_levels: dict[int, tuple[str, str]] = {}
    current_node_key: str | None = None
    current_code: str = ""

    checkpoints: list[Checkpoint] = []
    warnings: list[str] = []
    skipped_files_rows = 0
    skipped_context = 0
    optional_count = 0
    empty_code_rows = 0
    last_seq = 0

    for row_index in range(header_idx + 1, len(rows)):
        row = rows[row_index]
        raw_code = clean_display(_cell(row, "code"))
        raw_name = join_lines(_cell(row, "name"), "")
        archive = join_lines(_cell(row, "archive"), " ")
        files_text = _cell(row, "files")
        source = join_lines(_cell(row, "source"), " / ")
        filing = join_lines(_cell(row, "filing"), " / ")
        retention = {
            unit: join_lines(_cell(row, unit), " ").replace(" ", "")
            for unit in UNITS
        }

        if not any([raw_code, raw_name, archive, files_text, source, filing]):
            continue

        # ---------- 层级与类目节点 ----------
        if raw_code:
            if not raw_code.isdigit() or len(raw_code) > 4:
                warnings.append(
                    f"第 {row_index + 1} 行分类号“{raw_code}”不是 1~4 位数字，已按上一类目处理"
                )
                raw_code = ""

        if raw_code:
            level = len(raw_code)
            parent_key: str | None = None
            parent_code: str | None = None

            if level > 1:
                parent_level = level - 1
                parent_entry = current_levels.get(parent_level)
                if parent_entry is None:
                    # 没有紧邻的上级行时，退化为按编码前缀找已有节点
                    for node in reversed(list(nodes.values())):
                        if node.level == parent_level and raw_code.startswith(node.code):
                            parent_entry = (node.code, node.key)
                            break
                if parent_entry:
                    parent_code, parent_key = parent_entry
                    if not raw_code.startswith(parent_code):
                        warnings.append(
                            f"第 {row_index + 1} 行分类号 {raw_code} 与上级类目 {parent_code} "
                            f"前缀不一致，已按行序归入 {parent_code}，请核对模版"
                        )
                else:
                    warnings.append(
                        f"第 {row_index + 1} 行分类号 {raw_code} 找不到上级类目，已作为顶层节点处理"
                    )

            name = raw_name or (nodes[parent_key].name if parent_key and parent_key in nodes else "") or raw_code
            node_key = f"{parent_key or ''}|{raw_code}|{name}"
            if node_key not in nodes:
                nodes[node_key] = Node(
                    key=node_key,
                    code=raw_code,
                    name=name,
                    level=level,
                    parent_key=parent_key,
                    order=node_order,
                )
                node_order += 1

            current_levels[level] = (raw_code, node_key)
            for deeper in range(level + 1, 5):
                current_levels.pop(deeper, None)
            current_node_key = node_key
            current_code = raw_code
        else:
            empty_code_rows += 1
            if current_node_key is None:
                if files_text.strip():
                    warnings.append(f"第 {row_index + 1} 行没有分类号且无上文可继承，已跳过")
                continue
            if raw_name and raw_name != nodes[current_node_key].name:
                warnings.append(
                    f"第 {row_index + 1} 行分类号为空但类目名称“{raw_name}”与上文不同，"
                    f"已按上文“{nodes[current_node_key].name}”处理"
                )

        if current_node_key is None:
            continue
        node = nodes[current_node_key]

        # ---------- 主要项目文件 ----------
        if not files_text.strip():
            skipped_files_rows += 1
            continue

        bullets, context = split_bullets(files_text)
        skipped_context += len(context)
        if not bullets:
            skipped_files_rows += 1
            continue

        path_names = _build_path(nodes, current_node_key)
        for bullet, raw_line in bullets:
            seq = parse_seq(raw_line, last_seq + 1)
            last_seq = max(last_seq, seq)
            optional = is_optional_text(bullet)
            if optional:
                optional_count += 1

            cp_id = f"CP-{node.code}-{len(checkpoints) + 1:04d}"
            keywords = extract_terms(bullet)
            checkpoints.append(
                Checkpoint(
                    id=cp_id,
                    node_key=node.key,
                    code=node.code,
                    category=node.name,
                    path=path_names,
                    archive_content=archive or node.name,
                    text=bullet,
                    seq=seq,
                    source=source,
                    filing=filing,
                    retention=retention,
                    keywords=keywords,
                    optional=optional,
                    skip_reason="其他/备查类条目" if optional else "",
                    extra={"row": row_index + 1, "raw": raw_line},
                )
            )

    if not checkpoints:
        raise TemplateParseError("表结构识别成功，但没有解析到任何检查点，请检查“主要项目文件”列。")

    title = (doc.core_properties.title or "").strip()
    if not title:
        for paragraph in doc.paragraphs[:3]:
            if paragraph.text.strip():
                title = paragraph.text.strip()
                break
    title = title or path.stem

    doc_types = sorted({cp.code[0] for cp in checkpoints if cp.code})
    project_types = [
        {"code": node.code, "name": node.name}
        for node in sorted(nodes.values(), key=lambda item: item.order)
        if node.level == 2
    ]

    warnings.extend(
        [
            f"数据行：{len(rows) - header_idx - 1} 行；解析检查点 {len(checkpoints)} 条。",
            f"跳过无“主要项目文件”的行 {skipped_files_rows} 行；跳过引导句/范围说明 {skipped_context} 条。",
            f"分类号为空、按上一类目续行处理 {empty_code_rows} 行。",
            f"“其他/备查”类条目标记为不适用 {optional_count} 条。",
        ]
    )

    meta: dict[str, Any] = {
        "schema_version": 2,
        "source_file": path.name,
        "source_hash": _sha256(path),
        "template_name": title,
        "parsed_at": datetime.now().isoformat(timespec="seconds"),
        "doc_types": doc_types,
        "project_types": project_types,
        "total_checkpoints": len(checkpoints),
        "total_nodes": len(nodes),
        "table_rows": len(rows),
        "header_row": header_idx + 1,
        "columns": header,
    }

    return CheckpointSet(
        meta=meta,
        nodes=sorted(nodes.values(), key=lambda item: item.order),
        checkpoints=checkpoints,
        warnings=warnings,
    )
