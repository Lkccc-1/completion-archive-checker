"""Excel 整改单导出。"""
from __future__ import annotations

from datetime import datetime, timedelta
from io import BytesIO
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import STATE_MATCHED, STATE_MISSING, STATE_NA, STATE_SUSPECT, CheckReport
from .text import safe_excel_text

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=10)
TITLE_FONT = Font(name="微软雅黑", size=18, bold=True, color="1F4E79")
SECTION_FONT = Font(name="微软雅黑", size=12, bold=True, color="1F4E79")
NOTE_FONT = Font(name="微软雅黑", size=9, color="595959")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")

STATE_FILLS = {
    STATE_MISSING: PatternFill("solid", fgColor="FFC7CE"),
    STATE_SUSPECT: PatternFill("solid", fgColor="FFEB9C"),
    STATE_MATCHED: PatternFill("solid", fgColor="E2EFDA"),
    STATE_NA: PatternFill("solid", fgColor="EDEDED"),
    "整类缺失": PatternFill("solid", fgColor="F4B084"),
}
STATE_FONTS = {
    STATE_MISSING: Font(name="微软雅黑", size=10, color="9C0006"),
    STATE_SUSPECT: Font(name="微软雅黑", size=10, color="9C6500"),
    STATE_MATCHED: Font(name="微软雅黑", size=10, color="006100"),
    STATE_NA: Font(name="微软雅黑", size=10, color="595959"),
    "整类缺失": Font(name="微软雅黑", size=10, bold=True, color="833C0B"),
}


def _safe(value: Any) -> Any:
    if value is None:
        return ""
    return safe_excel_text(value)


def _write_row(ws, row_index: int, values: Sequence[Any], fill=None, font=None) -> None:
    for column, value in enumerate(values, start=1):
        cell = ws.cell(row=row_index, column=column, value=_safe(value))
        cell.font = font or BODY_FONT
        cell.alignment = WRAP
        cell.border = BORDER
        if fill:
            cell.fill = fill


def _write_header(ws, row_index: int, headers: Sequence[str]) -> None:
    for column, value in enumerate(headers, start=1):
        cell = ws.cell(row=row_index, column=column, value=_safe(value))
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def _set_widths(ws, widths: Sequence[float]) -> None:
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _table(
    ws,
    start_row: int,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    widths: Sequence[float],
    state_col: int | None = None,
) -> int:
    _write_header(ws, start_row, headers)
    _set_widths(ws, widths)
    row_index = start_row + 1
    for row in rows:
        _write_row(ws, row_index, row)
        if state_col is not None:
            state = str(row[state_col - 1])
            if state in STATE_FILLS:
                cell = ws.cell(row=row_index, column=state_col)
                cell.fill = STATE_FILLS[state]
                cell.font = STATE_FONTS[state]
        row_index += 1
    ws.freeze_panes = f"A{start_row + 1}"
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{max(start_row, row_index - 1)}"
    return row_index


def _page_setup(ws, landscape: bool = True) -> None:
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"
    ws.sheet_view.showGridLines = False


def _cover(wb: Workbook, report: CheckReport, project_name: str, inspector: str,
           reviewer: str, due_date: str) -> None:
    ws = wb.active
    ws.title = "封面"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 90

    ws.merge_cells("A1:B2")
    title = ws["A1"]
    title.value = "竣工资料完整性检查整改单"
    title.font = TITLE_FONT
    title.alignment = Alignment(horizontal="center", vertical="center")

    meta = report.meta
    stats = report.stats
    rows = [
        ("项目名称", project_name or meta.get("project_name", "")),
        ("资料类型", meta.get("doc_type_name", "")),
        ("项目类型", meta.get("project_type_name", "")),
        ("责任单位", meta.get("unit", "")),
        ("资料目录", meta.get("root", "")),
        ("检查标准", meta.get("template_name", "")),
        ("检查日期", datetime.now().strftime("%Y-%m-%d")),
        ("整改期限", due_date),
        ("检查人", inspector),
        ("复核人", reviewer),
        ("检查点总数", stats.get("total", 0)),
        ("已匹配", stats.get("matched", 0)),
        ("疑似", stats.get("suspect", 0)),
        ("缺失", stats.get("missing", 0)),
        ("不适用", stats.get("na", 0)),
        ("资料完整率", f"{stats.get('completeness', 0):.1%}"),
        ("未归类文件", len(report.orphan_files)),
    ]
    row_index = 4
    for label, value in rows:
        label_cell = ws.cell(row=row_index, column=1, value=label)
        label_cell.font = Font(name="微软雅黑", size=10, bold=True)
        label_cell.alignment = Alignment(vertical="center")
        label_cell.fill = PatternFill("solid", fgColor="F2F2F2")
        label_cell.border = BORDER
        value_cell = ws.cell(row=row_index, column=2, value=_safe(value))
        value_cell.font = BODY_FONT
        value_cell.alignment = Alignment(wrap_text=True, vertical="center")
        value_cell.border = BORDER
        row_index += 1

    row_index += 1
    note = ws.cell(
        row=row_index,
        column=1,
        value="说明：本表依据模版逐条比对文件名/文件夹路径生成，“疑似”项需人工确认后再整改；"
              "完整率 = 已匹配 /（检查点 - 不适用）。",
    )
    note.font = NOTE_FONT
    note.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=row_index, start_column=1, end_row=row_index + 1, end_column=2)

    row_index += 4
    sign = ws.cell(row=row_index, column=1, value="检查人签字：________________    复核人签字：________________    日期：________________")
    sign.font = BODY_FONT
    ws.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=2)

    _page_setup(ws, landscape=False)


def _rectification_sheet(wb: Workbook, report: CheckReport, due_date: str) -> None:
    ws = wb.create_sheet("整改单")
    headers = [
        "序号", "分类号", "类目名称", "归档内容", "主要项目文件", "来源", "立卷",
        "保管期限", "状态", "缺失说明", "责任单位", "整改期限", "整改要求", "依据", "备注",
    ]
    widths = [6, 10, 14, 14, 46, 12, 12, 16, 10, 32, 10, 12, 30, 24, 18]
    meta = report.meta
    unit = meta.get("unit", "")
    template = meta.get("template_name", "")
    rows: list[list[Any]] = []

    for gap in report.category_gaps:
        rows.append(
            [
                "",
                gap.code,
                gap.category,
                "—",
                f"整类缺失（{gap.count}条）",
                "",
                "",
                "",
                "整类缺失",
                "该分类下全部检查点均未发现对应资料",
                unit,
                due_date,
                "请按归档范围补齐该类全部资料",
                f"{template} · {gap.code}",
                " > ".join(gap.path),
            ]
        )

    missing = [item for item in report.results if item.state == STATE_MISSING]
    for index, item in enumerate(missing, start=1):
        cp = item.checkpoint
        rows.append(
            [
                index,
                cp.code,
                cp.category,
                cp.archive_content,
                cp.text,
                cp.source,
                cp.filing,
                cp.retention_text,
                item.state,
                item.reason,
                unit,
                due_date,
                f"请补充提供《{cp.text}》",
                f"{template} · {cp.code} {cp.category}",
                item.note,
            ]
        )

    if not rows:
        rows.append(["—"] * len(headers))
        rows[0][4] = "本次未发现缺项"
        rows[0][8] = "无"

    end = _table(ws, 1, headers, rows, widths, state_col=9)
    for row_index in range(2, end):
        state = str(ws.cell(row=row_index, column=9).value or "")
        if state == "整类缺失":
            for column in range(1, len(headers) + 1):
                ws.cell(row=row_index, column=column).fill = STATE_FILLS["整类缺失"]
    _page_setup(ws, landscape=True)


def _suspect_sheet(wb: Workbook, report: CheckReport) -> None:
    ws = wb.create_sheet("疑似清单")
    headers = [
        "序号", "分类号", "类目名称", "归档内容", "主要项目文件", "来源", "立卷",
        "保管期限", "置信度", "候选文件", "候选分数", "正文分", "印章", "说明", "备注",
    ]
    widths = [6, 10, 14, 14, 46, 12, 12, 16, 8, 50, 8, 10, 16, 30, 20]
    rows = []
    for index, item in enumerate(
        [item for item in report.results if item.state == STATE_SUSPECT], start=1
    ):
        cp = item.checkpoint
        first = item.evidence[0] if item.evidence else None
        rows.append(
            [
                index, cp.code, cp.category, cp.archive_content, cp.text, cp.source,
                cp.filing, cp.retention_text, item.confidence,
                first.rel_path if first else "", f"{first.score:.2f}" if first else "",
                item.content_text if item.content_pages else "", item.seal_text,
                item.reason, item.note,
            ]
        )
    if not rows:
        rows = [["—"] * len(headers)]
        rows[0][4] = "本次无疑似项"
        rows[0][8] = "—"
    end = _table(ws, 1, headers, rows, widths, state_col=None)
    _page_setup(ws, landscape=True)


def _matched_sheet(wb: Workbook, report: CheckReport) -> None:
    ws = wb.create_sheet("已匹配")
    headers = ["序号", "分类号", "类目名称", "归档内容", "主要项目文件", "状态", "置信度",
               "证据文件", "正文", "印章", "说明"]
    widths = [6, 10, 14, 14, 46, 10, 8, 52, 12, 16, 34]
    rows = []
    for index, item in enumerate(
        [item for item in report.results if item.state == STATE_MATCHED], start=1
    ):
        cp = item.checkpoint
        rows.append(
            [
                index, cp.code, cp.category, cp.archive_content, cp.text,
                item.state, item.confidence, item.evidence_text,
                item.content_text if item.content_pages else "", item.seal_text,
                item.reason,
            ]
        )
    if not rows:
        rows = [["—"] * len(headers)]
        rows[0][4] = "暂无已匹配项"
        rows[0][5] = "—"
    _table(ws, 1, headers, rows, widths, state_col=6)
    _page_setup(ws, landscape=True)


def _orphan_sheet(wb: Workbook, report: CheckReport) -> None:
    ws = wb.create_sheet("未归类文件")
    headers = ["序号", "相对路径", "扩展名", "大小(KB)", "最高分", "最佳候选检查点", "最佳候选分类号", "说明"]
    widths = [6, 70, 8, 10, 8, 46, 10, 34]
    rows = []
    for index, item in enumerate(report.orphan_files, start=1):
        rows.append(
            [
                index, item["rel_path"], item["ext"], round(item["size"] / 1024, 1),
                f"{item['best_score']:.2f}", item["best_cp_text"], item["best_cp_id"],
                item.get("note", ""),
            ]
        )
    if not rows:
        rows = [["—"] * len(headers)]
        rows[0][1] = "没有未归类文件"
    _table(ws, 1, headers, rows, widths)
    _page_setup(ws, landscape=True)


def _seal_sheet(wb: Workbook, report: CheckReport) -> None:
    """需要盖章的文种，逐条给出红色印章检测结论。"""
    ws = wb.create_sheet("印章核查")
    headers = ["序号", "分类号", "类目名称", "主要项目文件", "证据文件", "印章", "检出页码", "说明"]
    widths = [6, 10, 14, 44, 56, 14, 16, 34]
    rows = []
    for item in report.results:
        if not item.needs_seal:
            continue
        cp = item.checkpoint
        evidence = item.evidence[0].rel_path if item.evidence else ""
        if item.seal_state == "无":
            note = "该文种通常需盖章，未检出红色印章，请开卷确认"
        elif item.seal_state == "未知":
            note = "黑白扫描件，红色印章无法自动判定"
        elif item.seal_state == "有":
            note = "检出红色印章"
        else:
            note = "未读取文件内容"
        pages = "、".join(f"第{p}页" for p in item.seal_pages) or "—"
        rows.append([len(rows) + 1, cp.code, cp.category, cp.text, evidence,
                     item.seal_text, pages, note])
    if not rows:
        rows = [["—"] * len(headers)]
        rows[0][3] = "本次未识别到需要盖章的条目"
    _table(ws, 1, headers, rows, widths)
    _page_setup(ws, landscape=True)


def _stats_sheet(wb: Workbook, report: CheckReport) -> None:
    ws = wb.create_sheet("统计")
    widths = [14, 34, 10, 10, 10, 10, 10, 12]
    _set_widths(ws, widths)
    row = 1

    title = ws.cell(row=row, column=1, value="一、检查结果统计")
    title.font = SECTION_FONT
    row += 1
    total = max(1, report.stats.get("total", 0))
    rows = [
        ["状态", "数量", "占比"],
        [STATE_MATCHED, report.stats.get("matched", 0), f"{report.stats.get('matched', 0) / total:.1%}"],
        [STATE_SUSPECT, report.stats.get("suspect", 0), f"{report.stats.get('suspect', 0) / total:.1%}"],
        [STATE_MISSING, report.stats.get("missing", 0), f"{report.stats.get('missing', 0) / total:.1%}"],
        [STATE_NA, report.stats.get("na", 0), f"{report.stats.get('na', 0) / total:.1%}"],
        ["合计", report.stats.get("total", 0), "100.0%"],
        ["人工纠正", report.stats.get("manual", 0), ""],
        ["资料完整率", f"{report.stats.get('completeness', 0):.1%}", ""],
    ]
    for values in rows:
        _write_row(ws, row, values)
        row += 1
    for column in range(1, 4):
        cell = ws.cell(row=row - len(rows), column=column)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    row += 1
    title = ws.cell(row=row, column=1, value="二、按三级分类统计")
    title.font = SECTION_FONT
    row += 1
    headers = ["分类号", "类目名称", "检查点", "已匹配", "疑似", "缺失", "不适用", "完整率"]
    _write_row(ws, row, headers, fill=HEADER_FILL, font=HEADER_FONT)
    row += 1
    for item in report.stats.get("by_category", []):
        _write_row(
            ws,
            row,
            [
                item["code"], item["category"], item["total"], item["matched"],
                item["suspect"], item["missing"], item["na"], f"{item['completeness']:.1%}",
            ],
        )
        row += 1

    row += 1
    title = ws.cell(row=row, column=1, value="三、整类缺失")
    title.font = SECTION_FONT
    row += 1
    _write_row(ws, row, ["分类号", "类目名称", "缺失条数", "层级路径"], fill=HEADER_FILL, font=HEADER_FONT)
    row += 1
    for gap in report.category_gaps:
        _write_row(ws, row, [gap.code, gap.category, gap.count, " > ".join(gap.path)])
        row += 1
    if not report.category_gaps:
        _write_row(ws, row, ["—", "无", "", ""])
        row += 1

    row += 1
    title = ws.cell(row=row, column=1, value="四、筛选范围")
    title.font = SECTION_FONT
    row += 1
    meta = report.meta
    for label, value in [
        ("资料类型", meta.get("doc_type_name", "")),
        ("项目类型", meta.get("project_type_name", "")),
        ("责任单位", meta.get("unit", "")),
        ("检查点数量", meta.get("filtered_total", "")),
        ("立卷=相关条目", meta.get("related_count", 0)),
        ("匹配阈值", meta.get("match_threshold", "")),
        ("疑似阈值", meta.get("suspect_threshold", "")),
    ]:
        _write_row(ws, row, [label, value])
        row += 1

    _page_setup(ws, landscape=True)


def _notes_sheet(wb: Workbook, report: CheckReport) -> None:
    ws = wb.create_sheet("说明")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 120
    meta = report.meta
    content = meta.get("content") or {}
    if content.get("enabled"):
        method = (
            "文件名/路径匹配为主证据：文件夹中的 3~4 位分类号（如 8232）作为强信号，"
            "叠加关键词覆盖率与字符 3-gram 覆盖率。本期同时读取了 PDF/扫描件正文："
            f"共读 {content.get('files_read', 0)} 份文件、{content.get('ocr_pages', 0)} 页 OCR，"
            f"正文确认提升 {content.get('promoted', 0)} 条，正文救回缺失 {content.get('rescued', 0)} 条。"
        )
        limit = (
            f"正文只读前几页（读到能确认即停）。印章识别只认红色印章，"
            f"黑白扫描件无法判定（本次 {content.get('seals_unknown', 0)} 条不可判）；"
            "合卷 PDF 会表现为同一文件支撑多条，已在说明中标注“注意合卷核对”。"
        )
    else:
        method = (
            "按文件相对路径（文件夹名 + 文件名）与模版检查点比对。文件夹中的 3~4 位分类号"
            "（如 8232）作为强信号；同时使用关键词覆盖率和字符 3-gram 覆盖率计算置信度。"
        )
        limit = (
            "本期未读取文件正文，也未做印章识别。合卷 PDF（一份文件涵盖多个文件）"
            "会表现为同一文件支撑多条，已在说明中标注“注意合卷核对”。"
        )
    lines = [
        ("检查方法", method),
        ("判定口径", "已匹配：证据充分；疑似：有候选文件但证据不足，需人工确认；缺失：未发现候选；"
                    "不适用：模版中“其他/备查”类条目，未提供时不判缺。"),
        ("阈值", f"已匹配 ≥ {meta.get('match_threshold')}；疑似 ≥ {meta.get('suspect_threshold')}；"
                "分数低于疑似阈值判缺失。阈值可在后续版本中按项目校准。"),
        ("本期限制", limit),
        ("人工纠正", "疑似和缺项均可逐条改判为已提供/确认缺失/不适用；纠正结果按检查点编号落盘，"
                    "下次检查直接复用。"),
        ("模版", f"{meta.get('template_name', '')}（hash: {meta.get('template_hash', '')[:16]}…）"),
        ("生成时间", meta.get("generated_at", "")),
    ]
    row = 1
    for label, text in lines:
        label_cell = ws.cell(row=row, column=1, value=label)
        label_cell.font = Font(name="微软雅黑", size=10, bold=True, color="1F4E79")
        label_cell.alignment = Alignment(vertical="top")
        row += 1
        text_cell = ws.cell(row=row, column=1, value=text)
        text_cell.font = BODY_FONT
        text_cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 42
        row += 2
    _page_setup(ws, landscape=False)


def export_report(
    report: CheckReport,
    *,
    project_name: str = "",
    inspector: str = "",
    reviewer: str = "",
    due_days: int = 7,
) -> bytes:
    wb = Workbook()
    due_date = (datetime.now() + timedelta(days=max(1, due_days))).strftime("%Y-%m-%d")
    _cover(wb, report, project_name, inspector, reviewer, due_date)
    _rectification_sheet(wb, report, due_date)
    _suspect_sheet(wb, report)
    _matched_sheet(wb, report)
    _seal_sheet(wb, report)
    _orphan_sheet(wb, report)
    _stats_sheet(wb, report)
    _notes_sheet(wb, report)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
