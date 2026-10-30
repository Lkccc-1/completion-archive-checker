# 评测/评测3_标注表转CSV.py
# 把 标注表.xlsx（和可选的 补标表.xlsx）合并成第 3 步脚本要读的 标注结果.csv（utf-8-sig）
import csv, sys
from pathlib import Path
from openpyxl import load_workbook

根目录 = Path(__file__).resolve().parents[1]


def 绝对路径(值: str) -> Path:
    p = Path(值)
    return p if p.is_absolute() else 根目录 / p


# 用法：python 评测/评测3_标注表转CSV.py 主表.xlsx 标注结果.csv [补标表.xlsx]
主表 = 绝对路径(sys.argv[1]) if len(sys.argv) > 1 else 根目录 / "评测" / "标注表.xlsx"
出 = 绝对路径(sys.argv[2]) if len(sys.argv) > 2 else 根目录 / "评测" / "标注结果.csv"
# 补标表默认跟主表放同一目录，避免项目B误读项目A的补标表
补表 = 绝对路径(sys.argv[3]) if len(sys.argv) > 3 else 主表.parent / "补标表.xlsx"

认对 = {1, "1", "1.0", "√", "是", "对", "Y", "y"}
认错 = {0, "0", "0.0", "×", "否", "N", "n", "2"}


def 读一页(路径: Path, 页名: str) -> list[tuple]:
    if not 路径.exists():
        return []
    wb = load_workbook(路径, data_only=True)
    if 页名 not in wb.sheetnames:
        return []
    return [行[:6] for 行 in wb[页名].iter_rows(min_row=2, values_only=True) if 行 and 行[1]]


原始 = 读一页(主表, "标注")
if not 原始:
    raise SystemExit(f"读不到标注数据：{主表} 的「标注」页")
原始 += 读一页(补表, "补标")

行们, 怪值, 见过 = [], [], set()
组 = {}
for (序号, cid, 检查点, 候选, 标注, 备注) in 原始:
    if not cid or not 候选:
        continue
    键 = (str(cid).strip(), str(候选).strip())
    if 键 in 见过:
        continue
    见过.add(键)
    cid, 候选 = 键
    原 = 标注
    if isinstance(标注, str):
        标注 = 标注.strip()
    if 标注 in 认对:
        标注 = "1"
    elif 标注 in 认错 or 标注 in (None, ""):
        标注 = ""
    else:
        怪值.append((cid, 原))
        标注 = ""
    备注 = (备注 or "").strip() if isinstance(备注, str) else (备注 or "")
    行们.append({"cp_id": cid, "检查点": 检查点, "候选文件": 候选, "标注": 标注, "备注": 备注})
    组[cid] = 组.get(cid, 0) + (1 if 标注 == "1" else 0)

出.parent.mkdir(parents=True, exist_ok=True)
with open(出, "w", newline="", encoding="utf-8-sig") as fh:
    w = csv.DictWriter(fh, fieldnames=["cp_id", "检查点", "候选文件", "标注", "备注"])
    w.writeheader(); w.writerows(行们)

标了 = sum(1 for v in 组.values() if v)
没标 = [cid for cid, v in 组.items() if not v]
print(f"已生成 {出}：{len(行们)} 行 / {len(组)} 条检查点")
if 补表.exists():
    print(f"（已合并补标表：{补表.name}）")
print(f"填了 1 的检查点 {标了} 条，一条 1 都没填的 {len(没标)} 条")
if 怪值:
    print(f"注意：有 {len(怪值)} 个「标注」不是 1/0，已按空处理，例如：{怪值[:5]}")
if 没标:
    print("  没填 1 的前 10 条（确认确实没有归档，或补标）：")
    for cid in 没标[:10]:
        print("   ", cid)
