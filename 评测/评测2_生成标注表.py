# 评测/评测2_生成标注表.py
# 把 候选池.csv 排成方便人工标的 Excel：同一条检查点一个底色组、带序号、标注列只能填 1/0
import csv, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

根目录 = Path(__file__).resolve().parents[1]


def 绝对路径(值: str) -> Path:
    p = Path(值)
    return p if p.is_absolute() else 根目录 / p


# 用法：python 评测/评测2_生成标注表.py "评测/项目B/候选池.csv" "评测/项目B/标注表.xlsx"
源 = 绝对路径(sys.argv[1]) if len(sys.argv) > 1 else 根目录 / "评测" / "候选池.csv"
出 = 绝对路径(sys.argv[2]) if len(sys.argv) > 2 else 根目录 / "评测" / "标注表.xlsx"

rows = list(csv.DictReader(open(源, encoding="utf-8-sig")))
if not rows:
    raise SystemExit(f"候选池是空的：{源}")

wb = Workbook()
ws = wb.active
ws.title = "标注"

表头 = ["序号", "cp_id", "检查点", "候选文件", "标注", "备注"]
ws.append(表头)

组色 = PatternFill("solid", fgColor="EEF3FB")
白 = PatternFill("solid", fgColor="FFFFFF")
头色 = PatternFill("solid", fgColor="1F4FA8")
细分隔 = Side(style="thin", color="D7DFEA")
组分隔 = Side(style="medium", color="B9C9E0")

for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF", size=11)
    c.fill = 头色
    c.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 26

组号, 组内序号 = -1, 0
for r in rows:
    if r["cp_id"] != ws.cell(ws.max_row, 2).value:
        组号, 组内序号 = 组号 + 1, 0
    组内序号 += 1
    组首行 = 组内序号 == 1
    ws.append([组内序号, r["cp_id"], r["检查点"], r["候选文件"], "", r["备注"]])

    行 = ws.max_row
    底色 = 组色 if 组号 % 2 == 0 else 白
    for col in range(1, 7):
        cell = ws.cell(行, col)
        cell.font = Font(size=10)
        cell.fill = 底色
        cell.border = Border(top=(组分隔 if 组首行 else 细分隔), bottom=细分隔)
        cell.alignment = Alignment(
            horizontal="center" if col in (1, 5) else "left",
            vertical="top",
            wrap_text=col in (3, 4, 6),
        )

for 列, 宽 in zip("ABCDEF", [6, 14, 34, 96, 8, 14]):
    ws.column_dimensions[列].width = 宽

ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:F{ws.max_row}"

# 标注列只认 1 / 0，留空 = 不对（防止打成「是」「√」导致后面统计漏掉）
dv = DataValidation(type="list", formula1='"1,0"', allow_blank=True, showErrorMessage=True,
                    errorTitle="填 1 或 0", error="对填 1；不对留空或填 0。一条检查点可以有好几个 1。")
ws.add_data_validation(dv)
dv.add(f"E2:E{ws.max_row}")

说明 = wb.create_sheet("怎么标")
内容 = [
    ("怎么标（第 2 步）", ""),
    ("", ""),
    ("1", "同一 cp_id 的连续几行是同一组，隔组换了底色；序号是组内编号。"),
    ("2", "看「检查点」要的是哪份资料，在「候选文件」里找同文种 + 同主语的。"),
    ("3", "对的在「标注」填 1（有下拉框）；不对就留空。一条检查点可以有好几个 1。"),
    ("4", "一条都对不上：在这一组第一行的「备注」写 没有归档；拿不准写 存疑。"),
    ("5", "候选顺序是打乱的，不代表系统排序，别顺着第一行判。"),
    ("6", "标完按 Ctrl+S，再运行 评测3 转成 CSV（命令见本数据集目录下的 README）。"),
    ("", ""),
    ("判断规则（第 0 步，标注全程不变）", ""),
    ("同文种", "要「施工方案」就不能拿「施工方案报审表」或「检测方案」。"),
    ("同主语", "要「杆塔组立」就不能拿「基础施工」。"),
]
for 左, 右 in 内容:
    说明.append([左, 右])
    说明.cell(说明.max_row, 1).font = Font(bold=说明.max_row in (1, 10), size=14 if 说明.max_row == 1 else 11)
    说明.cell(说明.max_row, 2).alignment = Alignment(wrap_text=True, vertical="top")
说明.column_dimensions["A"].width = 16
说明.column_dimensions["B"].width = 82

出.parent.mkdir(parents=True, exist_ok=True)
wb.save(出)
print(f"已生成 {出}：{len(rows)} 行 / {组号 + 1} 组")
