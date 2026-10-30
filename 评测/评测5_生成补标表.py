# 评测/评测5_生成补标表.py
# 候选池有三路并集 + 上限 15，少数线上 top-1 仍可能没进池。
# 这个脚本把这些漏掉的 top-1 单独列成 补标表.xlsx，标完由 评测3 合并进 标注结果.csv。
import csv, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

根目录 = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(根目录))

from core.docx_parser import parse_docx
from core.scanner import scan_directory
from core.engine import filter_checkpoints
from core.matcher import Matcher

标准 = 根目录 / "电网建设项目、设备仪器类分类、归档范围及保管期限表.docx"


def 绝对路径(值: str) -> Path:
    p = Path(值)
    return p if p.is_absolute() else 根目录 / p


# 用法：python 评测/评测5_生成补标表.py "工程文件夹" "评测/项目B"
项目目录 = 绝对路径(sys.argv[1]) if len(sys.argv) > 1 else 根目录 / "（AX）项目A（220kV输电线路工程）扫描件"
数据目录 = 绝对路径(sys.argv[2]) if len(sys.argv) > 2 else 根目录 / "评测"
数据目录.mkdir(parents=True, exist_ok=True)
资料类型, 项目类型, 单位 = "8", "82", "施工"

cpset = parse_docx(标准)
files = scan_directory(项目目录)
cps = filter_checkpoints(cpset.checkpoints, 资料类型, 项目类型, 单位)

matcher = Matcher(cps, project_name=项目目录.name, known_codes=[c.code for c in cpset.checkpoints])
结果 = matcher.match(files)

池 = {}
for r in csv.DictReader(open(数据目录 / "候选池.csv", encoding="utf-8-sig")):
    池.setdefault(r["cp_id"], set()).add(r["候选文件"])

已标 = {}
标注结果 = 数据目录 / "标注结果.csv"
if 标注结果.exists():
    for r in csv.DictReader(open(标注结果, encoding="utf-8-sig")):
        if r["标注"].strip() == "1":
            已标.setdefault(r["cp_id"], set()).add(Path(r["候选文件"]).name)

要补 = []
for r in 结果:
    # 0 分 evidence 不是有效候选（个别条目会带一条空证据），列进补标表只会浪费标注时间
    if not r.evidence or r.score <= 0:
        continue
    top = r.evidence[0].rel_path
    cid = r.checkpoint.id
    if top in 池.get(cid, set()):
        continue
    要补.append((cid, r.checkpoint.text, top, r.score, sorted(已标.get(cid, set()))))

if not 要补:
    print("没有要补的，线上 top-1 都在候选池里")
    raise SystemExit(0)

wb = Workbook()
ws = wb.active
ws.title = "补标"
ws.append(["序号", "cp_id", "检查点", "候选文件", "标注", "备注", "该检查点已标为1的其它文件"])

头色 = PatternFill("solid", fgColor="B45309")
组色 = PatternFill("solid", fgColor="FDF0DD")
细 = Side(style="thin", color="E8D5B5")
粗 = Side(style="medium", color="D9A648")
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = 头色
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws.row_dimensions[1].height = 30

上一 = None
序号 = 0
for cid, text, top, sc, others in 要补:
    序号 = 序号 + 1 if cid == 上一 else 1
    上一 = cid
    ws.append([序号, cid, text, top, "", f"线上分 {sc:.2f}，原候选池没列到", " | ".join(others)])
    行 = ws.max_row
    for col in range(1, 8):
        cell = ws.cell(行, col)
        cell.font = Font(size=10)
        cell.fill = 组色
        cell.border = Border(top=(粗 if 序号 == 1 else 细), bottom=细)
        cell.alignment = Alignment(horizontal="center" if col in (1, 5) else "left",
                                   vertical="top", wrap_text=col in (3, 4, 6, 7))

for 列, 宽 in zip("ABCDEFG", [6, 14, 32, 90, 8, 22, 60]):
    ws.column_dimensions[列].width = 宽
ws.freeze_panes = "C2"
dv = DataValidation(type="list", formula1='"1,0"', allow_blank=True, showErrorMessage=True,
                    errorTitle="填 1 或 0", error="对填 1；不对留空或填 0。")
ws.add_data_validation(dv)
dv.add(f"E2:E{ws.max_row}")

说明 = wb.create_sheet("说明")
for 行 in [
    ["这一页是补的，行数取决于候选池覆盖情况"],
    [""],
    ["为什么有这一页", "候选池按上限截断、且只取三路并集，少数线上 top-1 可能没进池，单独列出来补标。"],
    ["怎么标", "和主表一样：对填 1，不对留空。最后一列是已经标过的其它正确文件，用来判断。"],
    ["标完", "保存后运行评测3（参数见数据集目录下的 README），它会自动合并这两张表。"],
]:
    说明.append(行)
说明.column_dimensions["A"].width = 18
说明.column_dimensions["B"].width = 90
说明["A1"].font = Font(bold=True, size=14)

出 = 数据目录 / "补标表.xlsx"
wb.save(出)
print(f"项目：{项目目录}")
print(f"已生成 {出}：{len(要补)} 行，涉及 {len({x[0] for x in 要补})} 条检查点")
if not 标注结果.exists():
    print("注意：还没读到 标注结果.csv，最后一列「已标为1的其它文件」会是空的；主表标完后再跑一次更好。")
